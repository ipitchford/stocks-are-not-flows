"""Source-bound extraction of the authorized open Gate A accounting subset."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime
import csv
import io
import math
from types import MappingProxyType
from typing import Iterable, Mapping, Sequence
import zipfile

import openpyxl

from housing_pressure.ident0.errors import InvariantViolation

from .config import GateAConfig
from .fetch import validate_source_payload
from .ods import read_ods_tables


@dataclass(frozen=True, slots=True)
class ValidationCheck:
    name: str
    status: str
    value: object
    expected: str
    detail: str


@dataclass(frozen=True, slots=True)
class DerivedTable:
    filename: str
    fieldnames: tuple[str, ...]
    rows: tuple[Mapping[str, object], ...]
    source_ids: tuple[str, ...]
    description: str


@dataclass(frozen=True, slots=True)
class GateAResult:
    schema_version: str
    decision: str
    full_gate_a_authorized: bool
    protocol_version: str
    core_period: str
    checks: tuple[ValidationCheck, ...]
    facts: Mapping[str, object]
    blocked_outputs: Mapping[str, str]
    tables: tuple[DerivedTable, ...]
    assurance_boundary: tuple[str, ...]

    def report_dict(self) -> dict[str, object]:
        return {
            "schema_version": self.schema_version,
            "decision": self.decision,
            "full_gate_a_authorized": self.full_gate_a_authorized,
            "protocol_version": self.protocol_version,
            "core_period": self.core_period,
            "checks": [asdict(item) for item in self.checks],
            "facts": dict(self.facts),
            "blocked_outputs": dict(self.blocked_outputs),
            "derived_tables": [
                {
                    "filename": item.filename,
                    "rows": len(item.rows),
                    "fields": list(item.fieldnames),
                    "source_ids": list(item.source_ids),
                    "description": item.description,
                }
                for item in self.tables
            ],
            "assurance_boundary": list(self.assurance_boundary),
        }


class _Checks:
    def __init__(self) -> None:
        self.items: list[ValidationCheck] = []

    def require(
        self,
        condition: bool,
        name: str,
        *,
        value: object,
        expected: str,
        detail: str,
    ) -> None:
        if not condition:
            raise InvariantViolation(
                f"Gate A validation failed at {name}: value={value!r}; expected {expected}"
            )
        self.items.append(
            ValidationCheck(
                name=name,
                status="PASS",
                value=value,
                expected=expected,
                detail=detail,
            )
        )


def _read_text(payload: bytes, source_id: str) -> str:
    try:
        return payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise InvariantViolation(f"source {source_id} is not UTF-8 CSV") from exc


def _float(value: object, label: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise InvariantViolation(f"{label} is not numeric: {value!r}") from exc
    if not math.isfinite(result):
        raise InvariantViolation(f"{label} is not finite")
    return result


def _row_by_label(rows: Sequence[Sequence[str]], label: str) -> Sequence[str]:
    matches = [row for row in rows if row and row[0] == label]
    if len(matches) != 1:
        raise InvariantViolation(f"expected one row labelled {label!r}, found {len(matches)}")
    return matches[0]


def _row_containing(rows: Sequence[Sequence[str]], value: str) -> int:
    matches = [index for index, row in enumerate(rows) if value in row]
    if len(matches) != 1:
        raise InvariantViolation(f"expected one row containing {value!r}, found {len(matches)}")
    return matches[0]


def _month_sequence(start: date, count: int) -> tuple[date, ...]:
    year, month = start.year, start.month
    values: list[date] = []
    for _ in range(count):
        values.append(date(year, month, 1))
        month += 1
        if month == 13:
            year += 1
            month = 1
    return tuple(values)


def _check_months(
    checks: _Checks,
    name: str,
    rows: Sequence[Mapping[str, object]],
    *,
    start: date,
    count: int,
) -> None:
    observed = tuple(row["date"] for row in rows)
    expected = _month_sequence(start, count)
    checks.require(
        observed == expected,
        name,
        value={
            "count": len(observed),
            "first": observed[0].isoformat() if observed else None,
            "last": observed[-1].isoformat() if observed else None,
        },
        expected=f"{count} unique consecutive months from {start.isoformat()}",
        detail="Missing or duplicate months fail closed; no interpolation is permitted.",
    )


def _extract_ehs(payload: bytes, checks: _Checks) -> tuple[dict[str, object], ...]:
    tables = read_ods_tables(payload)
    output: list[dict[str, object]] = []
    for start_year in range(2008, 2019):
        period = f"{start_year}-{str(start_year + 1)[-2:]}"
        if period not in tables:
            raise InvariantViolation(f"EHS workbook is missing sheet {period}")
        rows = tables[period]
        count_start = _row_containing(rows, "thousands of households") + 1
        count_end = _row_containing(rows, "percentages within tenure")
        count_rows = rows[count_start:count_end]
        owner = _row_by_label(count_rows, "all owner occupiers")
        social = _row_by_label(count_rows, "all social renters")
        private = _row_by_label(count_rows, "all private renters")
        total = _row_by_label(count_rows, "all tenures")

        within_age_start = _row_containing(rows, "percentages within age group") + 1
        within_age_rows = rows[within_age_start:]
        published_shares = {
            "owner": _row_by_label(within_age_rows, "all owner occupiers"),
            "social_renter": _row_by_label(within_age_rows, "all social renters"),
            "private_renter": _row_by_label(within_age_rows, "all private renters"),
        }

        def combined(row: Sequence[str], label: str) -> float:
            if len(row) <= 3:
                raise InvariantViolation(f"EHS {period} row {label} lacks age-band columns")
            return _float(row[2], f"EHS {period} {label} 25-34") + _float(
                row[3], f"EHS {period} {label} 35-44"
            )

        denominator = combined(total, "all tenures")
        counts = {
            "owner": combined(owner, "owner"),
            "social_renter": combined(social, "social renter"),
            "private_renter": combined(private, "private renter"),
        }
        shares = {name: 100.0 * value / denominator for name, value in counts.items()}
        count_residual = sum(counts.values()) - denominator
        share_residual = sum(shares.values()) - 100.0
        checks.require(
            abs(count_residual) <= 1e-6 and abs(share_residual) <= 1e-10,
            f"ehs_{period}_tenure_closure",
            value={"count_residual_thousands": count_residual, "share_residual_pp": share_residual},
            expected="three broad tenures close to the common 25-44 denominator",
            detail="The shares are recomputed from unrounded weighted counts, never summed percentages.",
        )
        row_map = {
            "owner": owner,
            "social_renter": social,
            "private_renter": private,
        }
        published_errors = [
            abs(
                100.0
                * _float(row_map[name][column], f"EHS {period} {name} count")
                / _float(total[column], f"EHS {period} all-tenure count")
                - _float(
                    published_shares[name][column],
                    f"EHS {period} published {name} percentage",
                )
            )
            for name in row_map
            for column in (2, 3)
        ]
        checks.require(
            max(published_errors) <= 1e-10,
            f"ehs_{period}_published_percentage_reproduction",
            value=max(published_errors),
            expected="maximum age-specific difference no greater than 1e-10 percentage points",
            detail="The count-based calculation reproduces the publisher's within-age table.",
        )
        sample_matches = [row for row in rows if row and row[0] == "sample size"]
        if len(sample_matches) > 1:
            raise InvariantViolation(f"EHS {period} has duplicate sample-size rows")
        if sample_matches:
            unweighted_sample: int | None = int(round(combined(sample_matches[0], "sample size")))
            sample_status = "published age-band sample sizes combined"
        else:
            header = rows[2] if len(rows) > 2 else ()
            if "sample size" not in header:
                raise InvariantViolation(f"EHS {period} exposes no sample-size information")
            unweighted_sample = None
            sample_status = "age-band unweighted sample unavailable; all-age total not substituted"
        output.append(
            {
                "period": period,
                "start_year": start_year,
                "geography": "England",
                "population": "private households with HRP age 25-44",
                "denominator_households_thousands": denominator,
                "unweighted_sample_size_25_44": unweighted_sample,
                "unweighted_sample_status": sample_status,
                "owner_households_thousands": counts["owner"],
                "owner_share_pct": shares["owner"],
                "private_renter_households_thousands": counts["private_renter"],
                "private_renter_share_pct": shares["private_renter"],
                "social_renter_households_thousands": counts["social_renter"],
                "social_renter_share_pct": shares["social_renter"],
                "uncertainty_status": "point estimate only; design covariance unavailable publicly",
                "source_id": "EHS_FA1201_OPEN",
            }
        )
    checks.require(
        len(output) == 11,
        "ehs_core_path",
        value=len(output),
        expected="11 EHS financial years from 2008-09 through 2018-19",
        detail="Post-2019 observations are excluded from the core.",
    )
    missing_sample_periods = [
        item["period"] for item in output if item["unweighted_sample_size_25_44"] is None
    ]
    checks.require(
        missing_sample_periods == ["2013-14"],
        "ehs_unweighted_sample_availability",
        value=missing_sample_periods,
        expected="only 2013-14 lacks a published 25-44 unweighted sample size",
        detail="The published all-age sample total is not mislabelled as the age-25-44 sample.",
    )
    return tuple(output)


def _extract_hpi(payload: bytes, checks: _Checks) -> tuple[dict[str, object], ...]:
    reader = csv.DictReader(io.StringIO(_read_text(payload, "HMLR_UKHPI_INDICES_2026_05")))
    expected_fields = ["Date", "Region_Name", "Area_Code", "Index"]
    checks.require(
        reader.fieldnames == expected_fields,
        "hpi_schema",
        value=reader.fieldnames,
        expected=str(expected_fields),
        detail="The compact attribute file has a different schema from the full UK HPI file.",
    )
    output: list[dict[str, object]] = []
    for row in reader:
        try:
            observed_date = datetime.strptime(row["Date"], "%Y-%m-%d").date()
        except ValueError as exc:
            raise InvariantViolation(f"invalid HPI date {row['Date']!r}") from exc
        if row["Area_Code"] == "E92000001" and date(2008, 1, 1) <= observed_date <= date(
            2019, 12, 1
        ):
            output.append(
                {
                    "date": observed_date,
                    "geography": row["Region_Name"],
                    "area_code": row["Area_Code"],
                    "index": _float(row["Index"], f"HPI {observed_date}"),
                    "source_id": "HMLR_UKHPI_INDICES_2026_05",
                }
            )
    output.sort(key=lambda item: item["date"])
    _check_months(checks, "hpi_england_months", output, start=date(2008, 1, 1), count=144)
    checks.require(
        all(item["geography"] == "England" and item["index"] > 0.0 for item in output),
        "hpi_geography_and_domain",
        value={"area_code": "E92000001", "positive": True},
        expected="England only with a positive unadjusted all-property index",
        detail="Buyer-status and financing cuts are not landlord identifiers.",
    )
    return tuple(output)


def _extract_iphrp(payload: bytes, checks: _Checks) -> tuple[dict[str, object], ...]:
    reader = csv.DictReader(io.StringIO(_read_text(payload, "ONS_IPHRP_HIST_V25")))
    required = {
        "v4_1",
        "Data Marking",
        "mmm-yy",
        "Time",
        "administrative-geography",
        "Geography",
        "index-and-year-change",
        "IndexAndYearChange",
    }
    checks.require(
        set(reader.fieldnames or ()) == required,
        "iphrp_schema",
        value=reader.fieldnames,
        expected="frozen version-25 long-format CSVW fields",
        detail="The observation value is v4_1; index and year-change rows share a file.",
    )
    output: list[dict[str, object]] = []
    for row in reader:
        if row["administrative-geography"] != "E92000001":
            continue
        if row["index-and-year-change"] != "index":
            continue
        try:
            observed_date = datetime.strptime(row["mmm-yy"], "%b-%y").date().replace(day=1)
        except ValueError as exc:
            raise InvariantViolation(f"invalid IPHRP month {row['mmm-yy']!r}") from exc
        if date(2008, 1, 1) <= observed_date <= date(2014, 12, 1):
            output.append(
                {
                    "date": observed_date,
                    "geography": row["Geography"],
                    "area_code": row["administrative-geography"],
                    "index": _float(row["v4_1"], f"IPHRP {observed_date}"),
                    "method_segment": "historical IPHRP v25; 2008-2014 only",
                    "source_id": "ONS_IPHRP_HIST_V25",
                }
            )
    output.sort(key=lambda item: item["date"])
    _check_months(checks, "iphrp_england_months", output, start=date(2008, 1, 1), count=84)
    checks.require(
        all(item["geography"] == "England" and float(item["index"]) > 0.0 for item in output),
        "iphrp_geography_and_domain",
        value={"geography": "England", "positive_index": True},
        expected="England only with a positive historical stock-rent index",
        detail="Area code and publisher geography name must agree.",
    )
    return tuple(output)


def _extract_pipr(payload: bytes, checks: _Checks) -> tuple[dict[str, object], ...]:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        raise InvariantViolation(f"cannot open PIPR XLSX: {exc}") from exc
    try:
        if "Table 1" not in workbook.sheetnames:
            raise InvariantViolation("PIPR workbook has no Table 1 sheet")
        worksheet = workbook["Table 1"]
        headers = tuple(cell.value for cell in next(worksheet.iter_rows(min_row=3, max_row=3)))
        required = ("Time period", "Area code", "Area name", "Index", "Rental price")
        checks.require(
            all(name in headers for name in required) and len(headers) == 40,
            "pipr_schema",
            value={
                "columns": len(headers),
                "required_present": all(name in headers for name in required),
            },
            expected="40-column Table 1 with time, area, overall index, and rental price",
            detail="The England region-or-country field is a marker and is not used as a filter.",
        )
        positions = {name: headers.index(name) for name in required}
        output: list[dict[str, object]] = []
        for values in worksheet.iter_rows(min_row=4, values_only=True):
            if values[positions["Area code"]] != "E92000001":
                continue
            raw_date = values[positions["Time period"]]
            if isinstance(raw_date, datetime):
                observed_date = raw_date.date()
            elif isinstance(raw_date, date):
                observed_date = raw_date
            else:
                continue
            if date(2015, 1, 1) <= observed_date <= date(2019, 12, 1):
                output.append(
                    {
                        "date": observed_date,
                        "geography": values[positions["Area name"]],
                        "area_code": values[positions["Area code"]],
                        "index": _float(values[positions["Index"]], f"PIPR {observed_date}"),
                        "estimated_monthly_rent_gbp": _float(
                            values[positions["Rental price"]], f"PIPR rent {observed_date}"
                        ),
                        "method_segment": "PIPR; 2015-2019 only",
                        "source_id": "ONS_PIPR_2026_07_22",
                    }
                )
    finally:
        workbook.close()
    output.sort(key=lambda item: item["date"])
    _check_months(checks, "pipr_england_months", output, start=date(2015, 1, 1), count=60)
    checks.require(
        all(
            item["geography"] == "England"
            and float(item["index"]) > 0.0
            and float(item["estimated_monthly_rent_gbp"]) > 0.0
            for item in output
        ),
        "pipr_geography_and_domain",
        value={"geography": "England", "positive_index_and_rent": True},
        expected="England only with positive PIPR index and estimated monthly rent",
        detail="Area code, publisher geography name, index, and rent must all agree.",
    )
    return tuple(output)


def _annual_hpi(
    hpi: Sequence[Mapping[str, object]], checks: _Checks
) -> tuple[dict[str, object], ...]:
    output: list[dict[str, object]] = []
    previous_december: float | None = None
    for year in range(2008, 2020):
        values = [float(item["index"]) for item in hpi if item["date"].year == year]
        if len(values) != 12:
            raise InvariantViolation(f"HPI year {year} does not contain 12 months")
        annual_mean = sum(values) / 12.0
        output.append(
            {
                "year": year,
                "geography": "England",
                "arithmetic_mean_index": annual_mean,
                "january_index": values[0],
                "december_index": values[-1],
                "december_to_december_100_log_change": (
                    None
                    if previous_december is None
                    else 100.0 * math.log(values[-1] / previous_december)
                ),
                "growth_support_status": (
                    "prior December outside extracted core"
                    if previous_december is None
                    else "complete December-to-December support"
                ),
                "source_id": "HMLR_UKHPI_INDICES_2026_05",
            }
        )
        previous_december = values[-1]
    base_mean = float(output[0]["arithmetic_mean_index"])
    for row in output:
        row["log_annual_mean_normalized_2008"] = math.log(
            float(row["arithmetic_mean_index"]) / base_mean
        )
    checks.require(
        len(output) == 12
        and output[0]["december_to_december_100_log_change"] is None
        and all(row["december_to_december_100_log_change"] is not None for row in output[1:]),
        "hpi_annual_support",
        value={"years": len(output), "growth_years": len(output) - 1},
        expected="12 annual means and 11 within-core December-to-December growth values",
        detail="The 2008 growth is missing because December 2007 is outside the extracted core.",
    )
    return tuple(output)


def _price_rent_ratio(
    hpi: Sequence[Mapping[str, object]],
    pipr: Sequence[Mapping[str, object]],
    checks: _Checks,
) -> tuple[dict[str, object], ...]:
    hpi_map = {item["date"]: float(item["index"]) for item in hpi}
    hpi_base = hpi_map[date(2015, 1, 1)]
    rent_base = float(pipr[0]["index"])
    output: list[dict[str, object]] = []
    for rent in pipr:
        observed_date = rent["date"]
        hpi_index = hpi_map.get(observed_date)
        if hpi_index is None:
            raise InvariantViolation(f"HPI has no month matching PIPR {observed_date}")
        rent_index = float(rent["index"])
        normalized_hpi = hpi_index / hpi_base
        normalized_rent = rent_index / rent_base
        price_to_rent = normalized_hpi / normalized_rent
        output.append(
            {
                "date": observed_date,
                "geography": "England",
                "hpi_normalized_jan2015": normalized_hpi,
                "stock_rent_index_normalized_jan2015": normalized_rent,
                "price_to_stock_rent_index": price_to_rent,
                "stock_rent_to_price_index": 1.0 / price_to_rent,
                "interpretation": "relative index only; not a gross or net rental yield",
                "source_ids": "HMLR_UKHPI_INDICES_2026_05|ONS_PIPR_2026_07_22",
            }
        )
    checks.require(
        len(output) == 60 and abs(float(output[0]["price_to_stock_rent_index"]) - 1.0) <= 1e-12,
        "price_rent_normalization",
        value={"months": len(output), "base": output[0]["price_to_stock_rent_index"]},
        expected="60 matched months and January 2015 ratio exactly one",
        detail="IPHRP and PIPR levels are never spliced; this ratio uses PIPR only.",
    )
    return tuple(output)


def _extract_supply(
    stock_payload: bytes,
    supply_payload: bytes,
    checks: _Checks,
) -> tuple[dict[str, object], ...]:
    stock_tables = read_ods_tables(stock_payload)
    if "LT_104" not in stock_tables:
        raise InvariantViolation("dwelling-stock workbook lacks LT_104")
    stock_rows = stock_tables["LT_104"]
    expected_stock_header = (
        "Date",
        "Year",
        "Owner occupied",
        "Rented privately or with a job or business",
        "Rented from private registered providers",
        "Rented from local authorities",
        "Other public sector dwellings",
        "All dwellings",
        "Notes",
    )
    if sum(row == expected_stock_header for row in stock_rows) != 1:
        raise InvariantViolation(
            "dwelling-stock workbook must contain the exact registered LT104 header"
        )
    stock: dict[int, float] = {}
    for row in stock_rows:
        if len(row) > 7 and row[1].isdigit():
            year = int(row[1])
            if 2008 <= year <= 2018:
                if row[0] != "31 March":
                    raise InvariantViolation(f"dwelling stock {year} is not dated 31 March")
                if year in stock:
                    raise InvariantViolation(f"dwelling stock contains duplicate year {year}")
                value = _float(row[7], f"dwelling stock {year}")
                if value <= 0.0:
                    raise InvariantViolation(f"dwelling stock {year} must be positive")
                stock[year] = value
    if set(stock) != set(range(2008, 2019)):
        raise InvariantViolation("dwelling stock must contain every core denominator year once")

    supply_tables = read_ods_tables(supply_payload)
    if "LT120_unrounded" not in supply_tables:
        raise InvariantViolation("net-supply workbook lacks LT120_unrounded")
    rows = supply_tables["LT120_unrounded"]
    header = next(
        (row for row in rows if row and row[0] == "Components of net housing supply"),
        None,
    )
    if header is None:
        raise InvariantViolation("net-supply workbook lacks its year header")
    labels = {
        name: _row_by_label(rows, name)
        for name in (
            "New build completions",
            "Net conversions",
            "Net change of use",
            "Net other gains",
            "Demolitions",
            "Census adjustments",
            "Total net additional dwellings",
        )
    }
    expected_periods = tuple(f"{year}-{str(year + 1)[-2:]}" for year in range(2008, 2019))
    positions: dict[str, int] = {}
    for period in expected_periods:
        matches = [index for index, value in enumerate(header) if value == period]
        if len(matches) != 1:
            raise InvariantViolation(f"net-supply header must contain exact period {period!r} once")
        positions[period] = matches[0]
    output: list[dict[str, object]] = []
    for period in expected_periods:
        column = positions[period]
        year = int(period[:4])
        values = {name: _float(row[column], f"{period} {name}") for name, row in labels.items()}
        component_sum = (
            values["New build completions"]
            + values["Net conversions"]
            + values["Net change of use"]
            + values["Net other gains"]
            - values["Demolitions"]
            + values["Census adjustments"]
        )
        total = values["Total net additional dwellings"]
        if total <= 0.0:
            raise InvariantViolation(f"{period} net additional dwellings must be positive")
        output.append(
            {
                "financial_year": period,
                "start_year": year,
                "geography": "England",
                "new_build_completions": values["New build completions"],
                "net_conversions": values["Net conversions"],
                "net_change_of_use": values["Net change of use"],
                "net_other_gains": values["Net other gains"],
                "demolitions": values["Demolitions"],
                "census_adjustments": values["Census adjustments"],
                "net_additional_dwellings": total,
                "published_component_residual": total - component_sum,
                "stock_reference_date": f"{year}-03-31",
                "start_of_year_all_dwellings": stock[year],
                "net_additions_rate_pct": 100.0 * total / stock[year],
                "denominator_note": "all dwellings including vacant dwellings",
                "source_ids": "MHCLG_DWELLING_STOCK_LT104_2026_05|MHCLG_NET_SUPPLY_LT120_2025_11",
            }
        )
    output.sort(key=lambda item: item["start_year"])
    checks.require(
        len(output) == 11 and [item["start_year"] for item in output] == list(range(2008, 2019)),
        "net_additions_core_years",
        value=[item["financial_year"] for item in output],
        expected="FY2008-09 through FY2018-19",
        detail="Each flow uses the all-dwelling stock at the start of its financial year.",
    )
    checks.require(
        max(abs(float(item["published_component_residual"])) for item in output) <= 50.0,
        "net_additions_component_reconciliation",
        value=max(abs(float(item["published_component_residual"])) for item in output),
        expected="absolute published component residual no greater than 50 dwellings",
        detail="Small source-table residuals are retained rather than silently forced to zero.",
    )
    return tuple(output)


def _extract_boe(payload: bytes, checks: _Checks) -> tuple[dict[str, object], ...]:
    reader = csv.DictReader(io.StringIO(_read_text(payload, "BOE_QUOTED_RATES_2008_2019")))
    expected_fields = ["DATE", "IUMBV34", "IUMZID4"]
    checks.require(
        reader.fieldnames == expected_fields,
        "boe_schema",
        value=reader.fieldnames,
        expected=str(expected_fields),
        detail="The query deliberately excludes daily Bank Rate to preserve monthly frequency.",
    )
    output: list[dict[str, object]] = []
    for row in reader:
        try:
            month_end = datetime.strptime(row["DATE"], "%d %b %Y").date()
        except ValueError as exc:
            raise InvariantViolation(f"invalid Bank date {row['DATE']!r}") from exc
        owner = _float(row["IUMBV34"], f"owner rate {month_end}")
        btl = None if not row["IUMZID4"] else _float(row["IUMZID4"], f"BTL rate {month_end}")
        output.append(
            {
                "date": date(month_end.year, month_end.month, 1),
                "publisher_date": month_end,
                "geography": "United Kingdom",
                "owner_2y_fixed_75ltv_pct": owner,
                "btl_2y_fixed_75ltv_pct": btl,
                "btl_minus_owner_spread_pp": None if btl is None else btl - owner,
                "interpretation": "quoted product rates; not realized borrower costs",
                "source_id": "BOE_QUOTED_RATES_2008_2019",
            }
        )
    output.sort(key=lambda item: item["date"])
    _check_months(checks, "boe_owner_months", output, start=date(2008, 1, 1), count=144)
    common = [item for item in output if item["btl_2y_fixed_75ltv_pct"] is not None]
    _check_months(checks, "boe_common_support", common, start=date(2012, 1, 1), count=96)
    checks.require(
        len(common) >= 36,
        "boe_registered_support_gate",
        value=len(common),
        expected="at least 36 consecutive common-support months",
        detail="BTL values before January 2012 remain missing and are never filled.",
    )
    checks.require(
        all(float(item["owner_2y_fixed_75ltv_pct"]) > 0.0 for item in output)
        and all(float(item["btl_2y_fixed_75ltv_pct"]) > 0.0 for item in common),
        "boe_positive_rate_domain",
        value={"owner_positive": True, "published_btl_positive": True},
        expected="all published quoted rates strictly positive",
        detail="Missing BTL months remain null and are excluded from the positivity check.",
    )
    return tuple(output)


def _annual_boe(
    rates: Sequence[Mapping[str, object]], checks: _Checks
) -> tuple[dict[str, object], ...]:
    output: list[dict[str, object]] = []
    for year in range(2008, 2020):
        annual = [item for item in rates if item["date"].year == year]
        if len(annual) != 12:
            raise InvariantViolation(f"Bank quoted-rate year {year} does not have 12 months")
        owner = [float(item["owner_2y_fixed_75ltv_pct"]) for item in annual]
        common = [item for item in annual if item["btl_2y_fixed_75ltv_pct"] is not None]
        output.append(
            {
                "year": year,
                "geography": "United Kingdom",
                "owner_support_months": len(owner),
                "owner_2y_fixed_75ltv_annual_mean_pct": sum(owner) / len(owner),
                "btl_common_support_months": len(common),
                "btl_2y_fixed_75ltv_annual_mean_pct": (
                    None
                    if not common
                    else sum(float(item["btl_2y_fixed_75ltv_pct"]) for item in common) / len(common)
                ),
                "btl_minus_owner_annual_mean_spread_pp": (
                    None
                    if not common
                    else sum(float(item["btl_minus_owner_spread_pp"]) for item in common)
                    / len(common)
                ),
                "common_support_status": (
                    "no published BTL observations"
                    if not common
                    else "complete 12-month common support"
                ),
                "interpretation": "arithmetic means of quoted product rates; not realized costs",
                "source_id": "BOE_QUOTED_RATES_2008_2019",
            }
        )
    observed_support = [int(item["btl_common_support_months"]) for item in output]
    checks.require(
        observed_support == [0, 0, 0, 0, 12, 12, 12, 12, 12, 12, 12, 12],
        "boe_annual_common_support",
        value=observed_support,
        expected="zero BTL months in 2008-2011 and twelve common months in every 2012-2019 year",
        detail="Annual BTL and spread values are missing unless all 12 months are available.",
    )
    return tuple(output)


def _table(
    filename: str,
    rows: Sequence[Mapping[str, object]],
    source_ids: Iterable[str],
    description: str,
) -> DerivedTable:
    if not rows:
        raise InvariantViolation(f"derived table {filename} has no rows")
    fieldnames = tuple(rows[0])
    if any(tuple(row) != fieldnames for row in rows):
        raise InvariantViolation(f"derived table {filename} has inconsistent fields")
    return DerivedTable(
        filename=filename,
        fieldnames=fieldnames,
        rows=tuple(MappingProxyType(dict(row)) for row in rows),
        source_ids=tuple(source_ids),
        description=description,
    )


def validate_payload_set(config: GateAConfig, payloads: Mapping[str, bytes]) -> None:
    """Bind source IDs to the exact registered bytes before any table is parsed."""

    expected_ids = {source.source_id for source in config.sources}
    supplied_ids = set(payloads)
    if supplied_ids != expected_ids:
        raise InvariantViolation(
            f"Gate A payload IDs differ; missing={sorted(expected_ids - supplied_ids)}, "
            f"unknown={sorted(supplied_ids - expected_ids)}"
        )
    for source in config.sources:
        validate_source_payload(source, payloads[source.source_id])


def extract_open_accounting(
    config: GateAConfig,
    payloads: Mapping[str, bytes],
) -> GateAResult:
    """Extract and validate the exact open subset authorized by Amendment 03."""

    config.validate()
    validate_payload_set(config, payloads)
    checks = _Checks()
    ehs = _extract_ehs(payloads["EHS_FA1201_OPEN"], checks)
    hpi = _extract_hpi(payloads["HMLR_UKHPI_INDICES_2026_05"], checks)
    iphrp = _extract_iphrp(payloads["ONS_IPHRP_HIST_V25"], checks)
    pipr = _extract_pipr(payloads["ONS_PIPR_2026_07_22"], checks)
    hpi_annual = _annual_hpi(hpi, checks)
    price_rent = _price_rent_ratio(hpi, pipr, checks)
    supply = _extract_supply(
        payloads["MHCLG_DWELLING_STOCK_LT104_2026_05"],
        payloads["MHCLG_NET_SUPPLY_LT120_2025_11"],
        checks,
    )
    rates = _extract_boe(payloads["BOE_QUOTED_RATES_2008_2019"], checks)
    rates_annual = _annual_boe(rates, checks)

    ehs_base, ehs_end = ehs[0], ehs[-1]
    hpi_base, hpi_end = hpi[0], hpi[-1]
    iphrp_base, iphrp_end = iphrp[0], iphrp[-1]
    pipr_base, pipr_end = pipr[0], pipr[-1]
    ratio_base, ratio_end = price_rent[0], price_rent[-1]
    common_rates = [item for item in rates if item["btl_2y_fixed_75ltv_pct"] is not None]
    facts: dict[str, object] = {
        "AF01_OWNER_25_44_POINT": {
            "geography": "England",
            "population": "private households with HRP age 25-44",
            "base_period": ehs_base["period"],
            "base_pct": ehs_base["owner_share_pct"],
            "end_period": ehs_end["period"],
            "end_pct": ehs_end["owner_share_pct"],
            "change_pp": float(ehs_end["owner_share_pct"]) - float(ehs_base["owner_share_pct"]),
            "uncertainty": "design covariance unavailable in open table",
            "viewing_status": "previously observed during feasibility",
        },
        "AF02_PRS_25_44_POINT": {
            "geography": "England",
            "population": "private households with HRP age 25-44",
            "base_period": ehs_base["period"],
            "base_pct": ehs_base["private_renter_share_pct"],
            "end_period": ehs_end["period"],
            "end_pct": ehs_end["private_renter_share_pct"],
            "change_pp": float(ehs_end["private_renter_share_pct"])
            - float(ehs_base["private_renter_share_pct"]),
            "uncertainty": "design covariance unavailable in open table",
            "viewing_status": "previously observed during feasibility",
        },
        "AF03_BROAD_TENURE_PATH": {
            "status": "point-estimate path reproduced for owner/private/social only",
            "years": len(ehs),
            "fine_tenure_and_covariance": "EUL blocked",
        },
        "AF07_HPI_PATH": {
            "geography": "England",
            "base_date": hpi_base["date"].isoformat(),
            "base_index": hpi_base["index"],
            "end_date": hpi_end["date"].isoformat(),
            "end_index": hpi_end["index"],
            "endpoint_pct_change": 100.0
            * (float(hpi_end["index"]) / float(hpi_base["index"]) - 1.0),
            "endpoint_100_log_change": 100.0
            * math.log(float(hpi_end["index"]) / float(hpi_base["index"])),
        },
        "AF08_RENT_SEGMENTS": {
            "geography": "England",
            "iphrp_period": f"{iphrp_base['date'].isoformat()} to {iphrp_end['date'].isoformat()}",
            "iphrp_pct_change": 100.0
            * (float(iphrp_end["index"]) / float(iphrp_base["index"]) - 1.0),
            "pipr_period": f"{pipr_base['date'].isoformat()} to {pipr_end['date'].isoformat()}",
            "pipr_pct_change": 100.0 * (float(pipr_end["index"]) / float(pipr_base["index"]) - 1.0),
            "level_splice": "prohibited",
        },
        "AF09_PRICE_RENT_INDEX": {
            "geography": "England",
            "period": f"{ratio_base['date'].isoformat()} to {ratio_end['date'].isoformat()}",
            "price_to_stock_rent_pct_change": 100.0
            * (float(ratio_end["price_to_stock_rent_index"]) - 1.0),
            "stock_rent_to_price_pct_change": 100.0
            * (float(ratio_end["stock_rent_to_price_index"]) - 1.0),
            "interpretation": "normalized relative index; not a gross or net rental yield",
            "viewing_status": "previously observed during feasibility",
        },
        "AF11_NET_ADDITIONS": {
            "geography": "England",
            "first_period": supply[0]["financial_year"],
            "first_rate_pct": supply[0]["net_additions_rate_pct"],
            "last_period": supply[-1]["financial_year"],
            "last_rate_pct": supply[-1]["net_additions_rate_pct"],
            "denominator": "all dwellings including vacant dwellings at start of FY",
        },
        "AF12_MORTGAGE_COST": {
            "geography": "United Kingdom",
            "common_start": common_rates[0]["date"].isoformat(),
            "common_end": common_rates[-1]["date"].isoformat(),
            "common_months": len(common_rates),
            "first_spread_pp": common_rates[0]["btl_minus_owner_spread_pp"],
            "last_spread_pp": common_rates[-1]["btl_minus_owner_spread_pp"],
            "mean_spread_pp": sum(float(item["btl_minus_owner_spread_pp"]) for item in common_rates)
            / len(common_rates),
            "annual_owner_path_years": len(rates_annual),
            "annual_btl_and_spread_complete_years": sum(
                int(item["btl_common_support_months"]) == 12 for item in rates_annual
            ),
            "interpretation": "quoted product rates; not realized borrower costs",
        },
    }
    tables = (
        _table(
            "ehs_tenure_25_44_england.csv",
            ehs,
            ("EHS_FA1201_OPEN",),
            "Open-table broad-tenure point estimates; no public design covariance.",
        ),
        _table(
            "hpi_monthly_england.csv",
            hpi,
            ("HMLR_UKHPI_INDICES_2026_05",),
            "Unadjusted all-property England HPI monthly path.",
        ),
        _table(
            "hpi_annual_england.csv",
            hpi_annual,
            ("HMLR_UKHPI_INDICES_2026_05",),
            "Annual mean, normalized log level, and December-to-December log growth.",
        ),
        _table(
            "rent_iphrp_2008_2014_england.csv",
            iphrp,
            ("ONS_IPHRP_HIST_V25",),
            "Historical stock-rent index segment; never level-spliced to PIPR.",
        ),
        _table(
            "rent_pipr_2015_2019_england.csv",
            pipr,
            ("ONS_PIPR_2026_07_22",),
            "PIPR stock-rent index segment and published estimated monthly rent.",
        ),
        _table(
            "price_to_stock_rent_2015_2019_england.csv",
            price_rent,
            ("HMLR_UKHPI_INDICES_2026_05", "ONS_PIPR_2026_07_22"),
            "January-2015-normalized relative index; not a rental yield.",
        ),
        _table(
            "net_additions_2008_2018_england.csv",
            supply,
            ("MHCLG_DWELLING_STOCK_LT104_2026_05", "MHCLG_NET_SUPPLY_LT120_2025_11"),
            "Financial-year net additions divided by start-of-year all-dwelling stock.",
        ),
        _table(
            "quoted_mortgage_rates_2008_2019_uk.csv",
            rates,
            ("BOE_QUOTED_RATES_2008_2019",),
            "Monthly quoted owner and BTL rates; missing BTL months remain blank.",
        ),
        _table(
            "quoted_mortgage_rates_annual_2008_2019_uk.csv",
            rates_annual,
            ("BOE_QUOTED_RATES_2008_2019",),
            "Annual arithmetic means; BTL and spread require all 12 common-support months.",
        ),
    )
    return GateAResult(
        schema_version="0.1",
        decision="OPEN-SUBSET-REPRODUCED+FULL-GATE-A-HELD",
        full_gate_a_authorized=False,
        protocol_version=config.protocol_version,
        core_period="England housing 2008-2019; GB/UK sources retain publisher geography",
        checks=tuple(checks.items),
        facts=MappingProxyType(facts),
        blocked_outputs=config.blocked_outputs,
        tables=tables,
        assurance_boundary=(
            "The open EHS results are point estimates without public design covariance.",
            "IPHRP and PIPR are separate method segments and are never level-spliced.",
            "The price-to-stock-rent index is not a gross or net rental yield.",
            "Dwelling-stock and household-tenure denominators are not interchangeable.",
            "Quoted mortgage rates are United Kingdom offers, not realized England borrower costs.",
            "No public series identifies the latent top non-housing resource flow.",
            "WAS net non-housing wealth outputs and full EHS inference remain EUL-blocked.",
            "This run cannot authorize structural estimation, unique attribution, or causal language.",
        ),
    )


__all__ = [
    "DerivedTable",
    "GateAResult",
    "ValidationCheck",
    "extract_open_accounting",
    "validate_payload_set",
]
